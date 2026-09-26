#!/usr/bin/env python3
"""Prepare four internal review sheets for the union of three retrievers' Top 20.

The workbook contains clinical text and must be generated and kept on Vaipe.
Each query-candidate pair appears once; a review credits every model that
selected that candidate. No similarity scores are shown to reviewers.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import random
import re
import socket

import numpy as np

from common import sha256, write_json
from moe_vote_rank import MODELS, QUERY_COUNT, TOP_K, rank_query, read_json, read_jsonl, validate_top20
from render_moe_vote_cases_html import checked_inputs


WORKBOOK_NAME = "moe_top20_union_review_4_people.xlsx"
SHEET_NAMES = ("Người 1", "Người 2", "Người 3", "Người 4")
SHEET_SIZES = (13, 13, 12, 12)
ALLOCATION_SEED = 20260926
FIRST_DATA_ROW = 11
RATINGS = (
    "0 - Không phù hợp", "1 - Ít phù hợp", "2 - Có liên quan",
    "3 - Rất phù hợp", "Chưa rõ",
)
DOCTOR_FLAGS = ("Có", "Không", "Chưa rõ")
MODEL_LABELS = {"fusion": "Retrieval", "openai": "OpenAI", "qwen3": "Qwen3"}
ILLEGAL_EXCEL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _source_path(manifest: dict, name: str) -> Path:
    source = manifest["sources"][name]
    path = Path(source["path"])
    if sha256(path) != source["sha256"]:
        raise ValueError(f"Source SHA-256 mismatch: {name}")
    return path


def _source_lists(results_root: Path, ranking_manifest: dict, selected: list[str]):
    """Recover the exact Top 20 used by the existing MoE ranking."""
    qwen_root = _source_path(ranking_manifest, "qwen_manifest").parent
    qwen_manifest = read_json(qwen_root / "manifest.json")
    for filename, expected in qwen_manifest["outputs"].items():
        if sha256(qwen_root / filename) != expected:
            raise ValueError(f"Qwen3 artifact SHA-256 mismatch: {filename}")
    ids = read_json(qwen_root / "patient_ids.json")
    if ids != sorted(ids) or len(ids) != len(set(ids)) or not set(selected) <= set(ids):
        raise ValueError("Invalid Qwen3 patient IDs")
    id_to_index = {value: index for index, value in enumerate(ids)}

    identity_path = _source_path(ranking_manifest, "identity_map")
    identity = list(read_jsonl(identity_path))
    raw_to_profile = {str(row["patient_id"]): row["profile_id"] for row in identity}
    if len(identity) != len(raw_to_profile) or set(raw_to_profile.values()) != set(ids):
        raise ValueError("Identity map is not one-to-one")
    profile_to_raw = {profile: raw for raw, profile in raw_to_profile.items()}

    import pandas as pd

    frame = pd.read_parquet(_source_path(ranking_manifest, "fusion_parquet"), columns=[
        "query_patient_id", "rank", "related_patient_id", "cosine_similarity",
    ])
    frame["query_patient_id"] = frame["query_patient_id"].astype(str)
    frame["related_patient_id"] = frame["related_patient_id"].astype(str)
    wanted_raw = {profile_to_raw[patient_id] for patient_id in selected}
    frame = frame[frame["query_patient_id"].isin(wanted_raw)]
    fusion = {}
    for raw_query, group in frame.groupby("query_patient_id"):
        query_id = raw_to_profile[raw_query]
        fusion[query_id] = [
            {"rank": int(row.rank), "patient_id": raw_to_profile[row.related_patient_id],
             "cosine_similarity": float(row.cosine_similarity)}
            for row in group.sort_values("rank").itertuples(index=False)
        ]
        validate_top20(query_id, fusion[query_id])

    openai = {}
    for row in read_jsonl(_source_path(ranking_manifest, "openai_top100")):
        query_id = row["query_patient_id"]
        if query_id in selected:
            if query_id in openai:
                raise ValueError("Duplicate OpenAI query")
            openai[query_id] = row["candidates"][:TOP_K]
            validate_top20(query_id, openai[query_id])

    queries = np.load(qwen_root / "query_embeddings.npy", allow_pickle=False)
    documents = np.load(qwen_root / "document_embeddings.npy", allow_pickle=False)
    if (queries.shape != documents.shape or queries.shape[0] != len(ids) or
            queries.dtype != np.float32 or documents.dtype != np.float32 or
            not np.isfinite(queries).all() or not np.isfinite(documents).all()):
        raise ValueError("Invalid Qwen3 vector shape or values")
    qwen = {}
    for query_id in selected:
        index = id_to_index[query_id]
        scores = documents @ queries[index]
        scores[index] = -np.inf
        order = np.argsort(-scores, kind="stable")[:TOP_K]
        qwen[query_id] = [
            {"rank": rank, "patient_id": ids[candidate_index],
             "cosine_similarity": float(scores[candidate_index])}
            for rank, candidate_index in enumerate(order.tolist(), 1)
        ]
        validate_top20(query_id, qwen[query_id])
    if set(fusion) != set(openai) or set(fusion) != set(qwen) or set(fusion) != set(selected):
        raise ValueError("Three-model query coverage differs")
    return {query_id: {"fusion": fusion[query_id], "openai": openai[query_id],
                       "qwen3": qwen[query_id]} for query_id in selected}


def _union_cases(results_root: Path, input_bundle: Path):
    ranking_rows, ranked_texts, ranking_manifest, input_path, input_manifest = checked_inputs(
        results_root, input_bundle
    )
    texts = {}
    for patient in read_jsonl(input_path):
        patient_id, value = patient["patient_id"], patient["embedding_text"]
        if (not re.fullmatch(r"P\d{5}", patient_id) or patient_id in texts or
                not isinstance(value, str) or not value or
                hashlib.sha256(value.encode("utf-8")).hexdigest() != patient["input_sha256"]):
            raise ValueError("Invalid full patient-text bundle")
        texts[patient_id] = value
    if len(texts) != input_manifest["patients"] or any(texts[key] != value for key, value in ranked_texts.items()):
        raise ValueError("Full patient-text bundle differs from ranked text")
    selected_path = results_root / "selected_queries.json"
    if sha256(selected_path) != ranking_manifest["outputs"][selected_path.name]:
        raise ValueError("Selected-query SHA-256 mismatch")
    selected = read_json(selected_path)
    if (len(selected) != len(set(selected)) or len(selected) != QUERY_COUNT or
            selected != [row["query_patient_id"] for row in ranking_rows]):
        raise ValueError("Selected queries differ from ranked queries")
    lists = _source_lists(results_root, ranking_manifest, selected)
    cases = {}
    for row in ranking_rows:
        query_id = row["query_patient_id"]
        sources = lists[query_id]
        expected_top20 = rank_query(query_id, sources)["top20"]
        if expected_top20 != row["top20"]:
            raise ValueError(f"MoE Top 20 changed for {query_id}")
        by_candidate = {}
        for model, candidates in sources.items():
            high = candidates[0]["cosine_similarity"]
            low = candidates[-1]["cosine_similarity"]
            for candidate in candidates:
                patient_id = candidate["patient_id"]
                score = candidate["cosine_similarity"]
                normalized = (score - low) / (high - low) if high > low else 0.5
                by_candidate.setdefault(patient_id, {})[model] = {
                    "rank": candidate["rank"], "normalized": normalized,
                }
        def order_key(item):
            patient_id, selected_by = item
            votes = len(selected_by)
            mean_rank = sum(value["rank"] for value in selected_by.values()) / votes
            one_vote_score = next(iter(selected_by.values()))["normalized"] if votes == 1 else 0.0
            return (-votes, mean_rank if votes > 1 else -one_vote_score, patient_id)
        ordered = sorted(by_candidate.items(), key=order_key)
        if (len(ordered) != row["unique_candidate_count"] or
                [patient_id for patient_id, _ in ordered[:TOP_K]] !=
                [item["patient_id"] for item in row["top20"]] or
                Counter(len(value) for value in by_candidate.values()) !=
                Counter({int(k): v for k, v in row["vote_counts_in_union"].items()})):
            raise ValueError(f"Union does not reconcile for {query_id}")
        moe_ranks = {item["patient_id"]: item["rank"] for item in row["top20"]}
        cases[query_id] = [{
            "query_id": query_id, "patient_id": patient_id,
            "moe_rank": moe_ranks.get(patient_id),
            "model_ranks": {model: selected_by[model]["rank"] for model in selected_by},
        } for patient_id, selected_by in ordered]
    return selected, cases, texts, ranking_manifest, input_path, input_manifest


def _allocate(selected: list[str], cases: dict[str, list[dict]]) -> list[list[str]]:
    """Keep 13/13/12/12 queries and minimize candidate-row imbalance."""
    rng = random.Random(ALLOCATION_SEED)
    best_score, best_groups = None, None
    for _ in range(30000):
        shuffled = selected.copy()
        rng.shuffle(shuffled)
        groups, offset = [], 0
        for size in SHEET_SIZES:
            groups.append(shuffled[offset:offset + size])
            offset += size
        loads = [sum(len(cases[patient_id]) for patient_id in group) for group in groups]
        score = (max(loads) - min(loads), sum((4 * value - sum(loads)) ** 2 for value in loads))
        if best_score is None or score < best_score:
            best_score, best_groups = score, groups
            if best_score[0] <= 1:
                break
    index = {patient_id: position for position, patient_id in enumerate(selected)}
    result = [sorted(group, key=index.__getitem__) for group in best_groups]
    if (list(map(len, result)) != list(SHEET_SIZES) or
            set(patient_id for group in result for patient_id in group) != set(selected)):
        raise ValueError("Invalid reviewer allocation")
    return result


def _plain_text(value: str) -> str:
    if len(value) > 32767 or ILLEGAL_EXCEL.search(value):
        raise ValueError("Clinical text cannot be represented unchanged in Excel")
    return value


def _write_workbook(path: Path, groups: list[list[str]], cases: dict, texts: dict):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.properties import CalcProperties

    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEET_NAMES:
        wb.create_sheet(name)
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    navy, blue = "18344B", "275D85"
    amber, cream = "FFF0BD", "FFF9E9"
    white, ink, green = "FFFFFF", "20313F", "DAF0E1"
    headers = ["Query ID", "Candidate ID", "MoE rank", "Query text", "Candidate text",
               "Retrieval rank", "OpenAI rank", "Qwen3 rank", "Mức phù hợp",
               "Nhận xét", "Bác sĩ xem?", "Hoàn tất"]
    model_columns = {"fusion": "F", "openai": "G", "qwen3": "H"}
    for sheet_index, (name, query_ids) in enumerate(zip(SHEET_NAMES, groups)):
        ws = wb[name]
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 85
        ws.sheet_properties.tabColor = ["275D85", "337B6D", "8A6839", "674D86"][sheet_index]
        ws.freeze_panes = "C11"
        widths = {"A": 14, "B": 16, "C": 11, "D": 74, "E": 74, "F": 14,
                  "G": 13, "H": 13, "I": 23, "J": 50, "K": 17, "L": 12}
        for column, width in widths.items():
            ws.column_dimensions[column].width = width
        ws.row_dimensions[1].height = 8
        ws.row_dimensions[2].height = 28
        ws.row_dimensions[3].height = 25
        ws.row_dimensions[4].height = 23
        ws.row_dimensions[5].height = 34
        ws.row_dimensions[9].height = 32
        ws.row_dimensions[10].height = 32
        ws["A2"] = f"Review candidate · {name}"
        ws["A2"].font = Font(name="Arial", size=14, bold=True, color=navy)
        ws["A3"], ws["C3"], ws["E3"], ws["G3"], ws["I3"] = (
            "Người phụ trách", "Query", "Candidate", "Đã hoàn tất", "Cả nhóm đã hoàn tất"
        )
        ws["B3"].fill = PatternFill("solid", fgColor=amber)
        ws["D3"] = len(query_ids)
        ws["F3"] = sum(len(cases[patient_id]) for patient_id in query_ids)
        ws["A4"] = "Nguồn: MoE 50 query và hồ sơ text Qwen3. Mỗi cặp chấm một lần; model có rank đều được tính."
        ws["I4"] = "MoE rank trống: ngoài Top 20 MoE."
        local_headers = ["Model", "Đã chọn", "Hoàn tất", "Phù hợp 2–3", "Chưa rõ", "Tỷ lệ 2–3/đã rõ"]
        for column, label in enumerate(local_headers, 1):
            ws.cell(5, column, label)
            ws.cell(5, column + 6, label)
        for row_number, model in enumerate(MODELS, 6):
            ws.cell(row_number, 1, MODEL_LABELS[model])
            ws.cell(row_number, 7, MODEL_LABELS[model])
        ws["A9"] = "Đánh giá: 0 không phù hợp; 1 ít phù hợp; 2 có liên quan; 3 rất phù hợp; Chưa rõ. Điền đánh giá, nhận xét và Bác sĩ xem? để hoàn tất."
        for column, label in enumerate(headers, 1):
            cell = ws.cell(10, column, label)
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(name="Arial", size=10, bold=True, color=white)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        data_rows = []
        for query_id in query_ids:
            for item in cases[query_id]:
                ranks = item["model_ranks"]
                data_rows.append([
                    query_id, item["patient_id"], item["moe_rank"],
                    _plain_text(texts[query_id]), _plain_text(texts[item["patient_id"]]),
                    ranks.get("fusion"), ranks.get("openai"), ranks.get("qwen3"),
                    None, None, None,
                ])
        end = 10 + len(data_rows)
        ws["H3"] = f'=COUNTIFS(L{FIRST_DATA_ROW}:L{end},1)'
        team_refs = ",".join(f"'{other}'!H3" for other in SHEET_NAMES)
        ws["L3"] = f"=SUM({team_refs})"
        for index, record in enumerate(data_rows, FIRST_DATA_ROW):
            for column, value in enumerate(record, 1):
                cell = ws.cell(index, column, value)
                cell.font = Font(name="Arial", size=10, color=ink)
                cell.alignment = Alignment(vertical="top", wrap_text=column in (4, 5, 10))
                if column in (1, 2, 4, 5) and isinstance(value, str):
                    cell.data_type = "s"  # Clinical text beginning '=' must not become a formula.
                if column in (9, 10, 11):
                    cell.fill = PatternFill("solid", fgColor=cream)
            ws.cell(index, 12, f'=IF(AND(I{index}<>"",LEN(TRIM(J{index}))>0,K{index}<>""),1,0)')
            ws.cell(index, 12).font = Font(name="Arial", size=10, color=ink)
            ws.cell(index, 12).alignment = Alignment(horizontal="center", vertical="top")
            if index == FIRST_DATA_ROW or data_rows[index - FIRST_DATA_ROW][0] != data_rows[index - FIRST_DATA_ROW - 1][0]:
                for cell in ws[index]:
                    cell.border = Border(top=Side(style="medium", color=blue))
                ws.cell(index, 1).font = Font(name="Arial", size=10, bold=True, color=blue)
            text_length = max(len(record[3]), len(record[4]))
            ws.row_dimensions[index].height = min(170, max(72, 18 * ((text_length + 95) // 96)))
        ws.auto_filter.ref = f"A10:L{end}"
        rating_validation = DataValidation(type="list", formula1='"' + ",".join(RATINGS) + '"', allow_blank=True)
        rating_validation.error = "Chọn mức phù hợp trong danh sách."
        rating_validation.errorTitle = "Mức phù hợp không hợp lệ"
        rating_validation.showErrorMessage = True
        ws.add_data_validation(rating_validation)
        rating_validation.add(f"I{FIRST_DATA_ROW}:I{end}")
        doctor_validation = DataValidation(type="list", formula1='"' + ",".join(DOCTOR_FLAGS) + '"', allow_blank=True)
        doctor_validation.error = "Chọn Có, Không hoặc Chưa rõ."
        doctor_validation.showErrorMessage = True
        ws.add_data_validation(doctor_validation)
        doctor_validation.add(f"K{FIRST_DATA_ROW}:K{end}")
        ws.conditional_formatting.add(
            f"I{FIRST_DATA_ROW}:I{end}",
            FormulaRule(formula=[f'OR(I{FIRST_DATA_ROW}="2 - Có liên quan",I{FIRST_DATA_ROW}="3 - Rất phù hợp")'],
                        fill=PatternFill("solid", fgColor=green)),
        )
        ws.conditional_formatting.add(
            f"L{FIRST_DATA_ROW}:L{end}",
            FormulaRule(formula=[f"L{FIRST_DATA_ROW}=1"], fill=PatternFill("solid", fgColor=green)),
        )
        for row_number, model in enumerate(MODELS, 6):
            model_range = f'{model_columns[model]}{FIRST_DATA_ROW}:{model_columns[model]}{end}'
            status_range = f'L{FIRST_DATA_ROW}:L{end}'
            rating_range = f'I{FIRST_DATA_ROW}:I{end}'
            ws.cell(row_number, 2, f'=COUNT({model_range})')
            ws.cell(row_number, 3, f'=COUNTIFS({model_range},">0",{status_range},1)')
            ws.cell(row_number, 4,
                    f'=COUNTIFS({model_range},">0",{status_range},1,{rating_range},"2 - Có liên quan")'
                    f'+COUNTIFS({model_range},">0",{status_range},1,{rating_range},"3 - Rất phù hợp")')
            ws.cell(row_number, 5, f'=COUNTIFS({model_range},">0",{status_range},1,{rating_range},"Chưa rõ")')
            ws.cell(row_number, 6, f'=IF(C{row_number}-E{row_number}=0,"",D{row_number}/(C{row_number}-E{row_number}))')
            for offset, local_column in enumerate("BCDE", 8):
                refs = ",".join(f"'{other}'!{local_column}{row_number}" for other in SHEET_NAMES)
                ws.cell(row_number, offset, f"=SUM({refs})")
            ws.cell(row_number, 12,
                    f'=IF(I{row_number}-K{row_number}=0,"",J{row_number}/(I{row_number}-K{row_number}))')
            for cell in ws[row_number][:12]:
                cell.font = Font(name="Arial", size=10, color=ink, bold=cell.column in (1, 7))
                cell.alignment = Alignment(horizontal="right" if cell.column not in (1, 7) else "left", vertical="center")
            ws.cell(row_number, 6).number_format = "0.0%"
            ws.cell(row_number, 12).number_format = "0.0%"
        for cell in ws[5][:12]:
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.font = Font(name="Arial", size=10, bold=True, color=white)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for address in ("A4", "I4", "A9"):
            ws[address].font = Font(name="Arial", size=10, italic=True, color="53697A")
        for cell in ws[3][:12]:
            if cell.coordinate != "B3":
                cell.font = Font(name="Arial", size=10, bold=cell.data_type != "f", color=ink)
        ws.sheet_format.defaultRowHeight = 20
        ws.print_title_rows = "1:10"
    wb.save(path)


def create(results_root: Path, input_bundle: Path, output: Path) -> None:
    if output.exists():
        raise FileExistsError(f"Choose a new output directory: {output}")
    selected, cases, texts, ranking_manifest, input_path, input_manifest = _union_cases(
        results_root, input_bundle
    )
    groups = _allocate(selected, cases)
    output.mkdir(parents=True)
    workbook_path = output / WORKBOOK_NAME
    _write_workbook(workbook_path, groups, cases, texts)
    query_counts = [len(group) for group in groups]
    row_counts = [sum(len(cases[query_id]) for query_id in group) for group in groups]
    manifest = {
        "schema_version": 1,
        "status": "prepared_for_internal_review_not_clinically_validated",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "host": socket.gethostname(),
        "raw_input": ranking_manifest["raw_input"],
        "raw_sha256": ranking_manifest["raw_sha256"],
        "ranking_manifest": str(results_root / "manifest.json"),
        "ranking_manifest_sha256": sha256(results_root / "manifest.json"),
        "patient_content_path": str(input_path),
        "patient_content_sha256": input_manifest["output_sha256"],
        "code_sha256": sha256(Path(__file__)),
        "candidate_scope": "Union of each model's Top 20 for each of the 50 queries",
        "assignment_seed": ALLOCATION_SEED,
        "query_count": QUERY_COUNT,
        "query_counts": query_counts,
        "pair_count": sum(row_counts),
        "pair_counts": row_counts,
        "model_selections": {model: QUERY_COUNT * TOP_K for model in MODELS},
        "contains_clinical_text": True,
        "similarity_scores_displayed": False,
        "rating_values": list(RATINGS),
        "positive_ratings": list(RATINGS[2:4]),
        "sheets": {name: group for name, group in zip(SHEET_NAMES, groups)},
        "outputs": {WORKBOOK_NAME: sha256(workbook_path)},
    }
    write_json(output / "manifest.json", manifest)
    verify(output)
    print(json.dumps({"workbook": str(workbook_path), "queries": query_counts,
                      "candidate_pairs": row_counts, "total_pairs": sum(row_counts)}, ensure_ascii=False))


def verify(output: Path) -> None:
    from openpyxl import load_workbook

    manifest = read_json(output / "manifest.json")
    workbook_path = output / WORKBOOK_NAME
    if (manifest["status"] != "prepared_for_internal_review_not_clinically_validated" or
            manifest["query_count"] != QUERY_COUNT or
            manifest["contains_clinical_text"] is not True or
            manifest["similarity_scores_displayed"] is not False or
            sha256(workbook_path) != manifest["outputs"][WORKBOOK_NAME] or
            sha256(Path(manifest["raw_input"])) != manifest["raw_sha256"] or
            sha256(Path(manifest["ranking_manifest"])) != manifest["ranking_manifest_sha256"] or
            sha256(Path(manifest["patient_content_path"])) != manifest["patient_content_sha256"]):
        raise ValueError("Workbook manifest or source SHA-256 mismatch")
    ranking_root = Path(manifest["ranking_manifest"]).parent
    input_bundle = Path(manifest["patient_content_path"]).parent
    selected, cases, texts, _, _, _ = _union_cases(ranking_root, input_bundle)
    wb = load_workbook(workbook_path, read_only=False, data_only=False)
    if (wb.sheetnames != list(SHEET_NAMES) or wb.calculation.calcMode != "auto" or
            not wb.calculation.fullCalcOnLoad):
        raise ValueError("Workbook sheet or recalculation contract failed")
    seen_queries, pair_count = [], 0
    for sheet_index, name in enumerate(SHEET_NAMES):
        ws = wb[name]
        assigned = manifest["sheets"][name]
        expected = [item for query_id in assigned for item in cases[query_id]]
        end = 10 + len(expected)
        if (len(assigned) != SHEET_SIZES[sheet_index] or
                len(expected) != manifest["pair_counts"][sheet_index] or
                ws.max_row != end or ws.freeze_panes != "C11" or
                ws.auto_filter.ref != f"A10:L{end}" or len(ws.data_validations.dataValidation) != 2):
            raise ValueError(f"Sheet structure failed: {name}")
        seen_queries.extend(assigned)
        pair_count += len(expected)
        by_query_model = Counter()
        for row_index, expected_item in enumerate(expected, FIRST_DATA_ROW):
            query_id, patient_id = expected_item["query_id"], expected_item["patient_id"]
            actual = [ws.cell(row_index, column).value for column in range(1, 13)]
            source_ranks = expected_item["model_ranks"]
            if (actual[:3] != [query_id, patient_id, expected_item["moe_rank"]] or
                    actual[3] != texts[query_id] or actual[4] != texts[patient_id] or
                    actual[5:8] != [source_ranks.get(model) for model in MODELS] or
                    actual[8:11] != [None, None, None] or
                    actual[11] != f'=IF(AND(I{row_index}<>"",LEN(TRIM(J{row_index}))>0,K{row_index}<>""),1,0)'):
                raise ValueError(f"Wrong or missing review row: {name}:{row_index}")
            for model in source_ranks:
                by_query_model[(query_id, model)] += 1
        if any(by_query_model[(query_id, model)] != TOP_K for query_id in assigned for model in MODELS):
            raise ValueError(f"Incomplete model Top 20: {name}")
        for row_number, model in enumerate(MODELS, 6):
            if (ws.cell(row_number, 2).data_type != "f" or
                    ws.cell(row_number, 3).data_type != "f" or
                    ws.cell(row_number, 4).data_type != "f" or
                    ws.cell(row_number, 5).data_type != "f" or
                    ws.cell(row_number, 6).data_type != "f" or
                    ws.cell(row_number, 8).data_type != "f" or
                    ws.cell(row_number, 12).data_type != "f"):
                raise ValueError(f"Missing summary formula: {name}:{model}")
        if ws["H3"].data_type != "f" or ws["L3"].data_type != "f":
            raise ValueError(f"Missing progress formula: {name}")
    if (len(seen_queries) != QUERY_COUNT or len(set(seen_queries)) != QUERY_COUNT or
            set(seen_queries) != set(selected) or pair_count != manifest["pair_count"] or
            pair_count != sum(len(value) for value in cases.values())):
        raise ValueError("Workbook query or pair coverage failed")
    print(json.dumps({"status": "pass", "sheets": list(SHEET_NAMES),
                      "queries": QUERY_COUNT, "candidate_pairs": pair_count,
                      "model_selections_each": QUERY_COUNT * TOP_K}, ensure_ascii=False))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    make = sub.add_parser("create")
    make.add_argument("--results-root", type=Path, required=True)
    make.add_argument("--input-bundle", type=Path, required=True)
    make.add_argument("--output", type=Path, required=True)
    check = sub.add_parser("verify")
    check.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.command == "create":
        create(args.results_root.resolve(), args.input_bundle.resolve(), args.output.resolve())
    else:
        verify(args.output.resolve())


if __name__ == "__main__":
    main()
