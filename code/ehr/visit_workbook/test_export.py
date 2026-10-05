"""Meaningful lineage and Excel-boundary checks; execute only on Vaipe."""
from contextlib import redirect_stdout
from datetime import datetime
import io
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
import export as e


class ExportTests(unittest.TestCase):
    def test_surgery_ambiguity_including_outside_cohort_is_not_resolved(self):
        cohort = {'a'}
        mapping = {'good': {'a'}, 'ambiguous': {'a', 'outside'}}
        self.assertEqual(e.resolve_link(e.SURGERY, '', 'good', mapping, cohort), ('a', '', ['a']))
        self.assertEqual(e.resolve_link(e.SURGERY, '', 'ambiguous', mapping, cohort),
                         ('', 'ma_chi_dinh_khong_duy_nhat', ['a', 'outside']))
        self.assertEqual(e.resolve_link(e.SURGERY, '', '', {'': {'a'}}, cohort)[1], 'khong_noi_duoc_ma_chi_dinh')

    def test_utf16_continuations_reassemble_unchanged(self):
        original = 'á' * 29999 + '😀' * 20000 + '=literal'
        parts = e.split_text(original)
        self.assertEqual(''.join(parts), original)
        self.assertTrue(all(e.utf16len(p) <= e.LIMIT for p in parts))
        self.assertEqual(e.unique_headers(['BenhAn_Id', 'BenhAn_Id']), ['BenhAn_Id', 'BenhAn_Id::__cot_2'])

    def test_full_export_preserves_repeated_results_drugs_and_exceptions(self):
        with tempfile.TemporaryDirectory() as folder:
            raw = Path(folder) / 'raw.xlsx'
            run = Path(folder) / 'output'
            wb = Workbook()
            wb.remove(wb.active)
            main = wb.create_sheet(e.MAIN)
            main.append(['SoBenhAn', 'SoVaoVien', 'NgayVaoVien', 'BenhAn_Id', 'BenhAn_Id', 'QuaTrinhBenhLy'])
            main.append(['0001', '0009', datetime(2020, 1, 1), 'first', 'second', '=literal'])
            orders = wb.create_sheet(e.ORDERS)
            orders.append(['SoBenhAn', 'SoVaoVien', 'YeuCauChiTiet_Id', 'TenBenhNhan', 'NamSinh', 'GioiTinh'])
            orders.append(['0001', '0009', 'o', 'Test', 1980, 1])
            orders.append(['outside', 'outsider', 'ambiguous', 'Other', 1981, 2])
            orders.append(['0001', '0009', 'ambiguous', 'Test', 1980, 1])
            meds = wb.create_sheet(e.MEDS)
            meds.append(['sobenhan', 'TenBenhNhan', 'GioiTinh', 'TenDuoc', 'GhiChu'])
            meds.append(['0001', 'Test', 'nam', 'Drug A', 'same'])
            meds.append(['0001', 'Test', 'nam', 'Drug A', 'same'])
            results = wb.create_sheet(e.RESULTS)
            results.append(['SoBenhAn', 'YeuCauChiTiet_Id', 'TEN_CHI_SO', 'GIA_TRI', 'MO_TA'])
            results.append(['0001', 'o', 'Lab', '<5', 'x' * 31000 + '😀' * 5000])
            results.append(['0001', 'o', 'Lab', None, 'NULL'])
            surgery = wb.create_sheet(e.SURGERY)
            surgery.append(['YeuCauChiTiet_Id', 'ma_lk', 'TenDichVu'])
            surgery.append(['o', 'uninterpreted', 'Procedure'])
            surgery.append(['ambiguous', 'uninterpreted', 'Unresolved'])
            wb.save(raw)
            with redirect_stdout(io.StringIO()):
                e.prepare(run, raw)
                e.export(run, 1000000)
            manifest = json.loads((run / 'manifest.json').read_text())
            self.assertEqual(manifest['rows'], 1)
            self.assertEqual(manifest['exception_rows'], 2)
            self.assertEqual(manifest['source_membership_verified'][e.MEDS], 2)
            saved = load_workbook(manifest['output_path'], read_only=True, data_only=False)
            iterator = saved['Visits'].iter_rows(values_only=True)
            row = dict(zip(next(iterator), next(iterator)))
            self.assertEqual(row['SoBenhAn'], '0001')
            self.assertEqual(row['SoVaoVien'], '0009')
            self.assertEqual(row['QuaTrinhBenhLy'], '=literal')
            self.assertEqual(row['BenhAn_Id'], 'first')
            self.assertEqual(row['BenhAn_Id::__cot_2'], 'second')
            self.assertEqual(row['SoKetQuaKQCLS'], 2)
            self.assertEqual(row['SoPhauThuatThuThuat'], 1)
            self.assertEqual(row['GioiTinh_ma_DVKT'], 1)
            self.assertEqual(row['GioiTinh_Thuoc'], 'nam')
            saved.close()


if __name__ == '__main__':
    unittest.main()
