"""Render a PHI-free layout preview of the saved XLSX on Vaipe, using Pillow.

All visit and exception data cells are replaced, never copied into the image.
Only actual headers and the nonclinical schema/reading guide are displayed.
This is a layout inspection, not a substitute for exact server data checks.
"""
import json
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


def main(path):
    wb = load_workbook(path, read_only=True, data_only=False)
    font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 14)
    windows = [('Visits', 1, 7, 4), ('Visits', 8, 12, 4), ('Cau_truc', 1, 3, 10)]
    if 'Ngoai_le' in wb.sheetnames:
        windows.append(('Ngoai_le', 1, 5, 3))
    panels = []
    for sheetname, startcol, endcol, maxrow in windows:
        ws = wb[sheetname]
        column_widths, row_heights = {}, {}
        with ZipFile(path) as archive:
            with archive.open(f'xl/worksheets/sheet{wb.sheetnames.index(sheetname)+1}.xml') as stream:
                for event, element in ET.iterparse(stream, events=('start', 'end')):
                    tag = element.tag.rsplit('}', 1)[-1]
                    if event == 'start' and tag == 'col':
                        for c in range(int(element.attrib['min']), int(element.attrib['max'])+1):
                            column_widths[c] = float(element.attrib.get('width', 13))
                    if event == 'start' and tag == 'row':
                        r = int(element.attrib['r'])
                        row_heights[r] = float(element.attrib.get('ht', 15))
                        if r > maxrow:
                            break
                    if event == 'end':
                        element.clear()
        widths = [int(column_widths.get(c, 13) * 7 + 8) for c in range(startcol, endcol+1)]
        heights = [int(row_heights.get(r, 15) * 4/3) for r in range(1, maxrow+1)]
        source_rows = list(ws.iter_rows(max_row=maxrow, max_col=endcol))
        panel = Image.new('RGB', (sum(widths)+1, sum(heights)+40), 'white')
        d = ImageDraw.Draw(panel)
        d.text((8, 8), f'{sheetname}: {get_column_letter(startcol)}–{get_column_letter(endcol)} (data redacted)', font=font, fill='black')
        y = 38
        for r, height in enumerate(heights, 1):
            x = 0
            for c, width in zip(range(startcol, endcol+1), widths):
                source = source_rows[r-1][c-1]
                value = source.value
                if sheetname != 'Cau_truc' and r > 1:
                    value = '[data]' if value is not None else ''
                text = '' if value is None else str(value)
                tile = Image.new('RGB', (width-1, height-1), 'white')
                td = ImageDraw.Draw(tile)
                lines, current = [], ''
                if getattr(source.alignment, 'wrap_text', False):
                    for ch in text:
                        if ch == '\n' or td.textlength(current+ch, font=font) > width-8:
                            lines.append(current)
                            current = '' if ch == '\n' else ch
                        else:
                            current += ch
                    lines.append(current)
                else:
                    lines = [text]
                td.multiline_text((3, 3), '\n'.join(lines), font=font, fill='black', spacing=2)
                panel.paste(tile, (x+1, y+1))
                d.rectangle((x, y, x+width, y+height), outline='#d0d0d0')
                x += width
            y += height
        panels.append(panel)
    result = Image.new('RGB', (max(p.width for p in panels), sum(p.height for p in panels)+16*(len(panels)-1)), 'white')
    y = 0
    for p in panels:
        result.paste(p, (0, y))
        y += p.height+16
    destination = path.parent / 'layout_redacted.png'
    result.save(destination)
    wb.close()
    print(json.dumps({'preview_path': str(destination), 'clinical_values_redacted': True}))


if __name__ == '__main__':
    main(Path(sys.argv[1]))
